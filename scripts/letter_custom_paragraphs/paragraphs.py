#!/usr/bin/env python3
"""
Generate personalized letter paragraphs from agency websites.

Per recipient:
  1. Fetch homepage, pick relevant subpages (services, industries, case studies ...)
  2. Clean HTML (nav, footer, cookie banners, scripts), drop duplicate lines
  3. Send text to a model (Claude API or local via LM Studio / Ollama)
  4. Validate JSON response (word count, exclamation marks, agency name ...) and save

Commands:
  python paragraphs.py crawl   list.csv [--limit N] [--refresh]
  python paragraphs.py run     list.csv --profile sonnet --profile qwen-ollama [--limit 10] [--batch]
  python paragraphs.py prompt  list.csv --row 1     (print the full prompt for one recipient)
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import random
import re
import sys
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path
from typing import NamedTuple
from urllib.parse import urldefrag, urljoin, urlparse

import requests
import yaml
from bs4 import BeautifulSoup

try:
    import trafilatura
except ImportError:  # fallback works without it
    trafilatura = None

BASE_DIR = Path(__file__).resolve().parent
CACHE_DIR = BASE_DIR / "cache" / "websites"
RUNS_DIR = BASE_DIR / "runs"

USER_AGENT = ("Mozilla/5.0 (Macintosh; Intel Mac OS X 14_5) AppleWebKit/537.36 "
              "(KHTML, like Gecko) Chrome/126.0 Safari/537.36")
FILE_EXTENSIONS = (".pdf", ".jpg", ".jpeg", ".png", ".gif", ".svg", ".webp", ".zip",
                   ".mp4", ".mp3", ".doc", ".docx", ".xls", ".xlsx", ".ics", ".xml")
NOISE = re.compile(
    r"(cookie|akzeptieren|alle annehmen|einstellungen speichern|datenschutzerklärung|"
    r"javascript|©|copyright|alle rechte vorbehalten|zum inhalt springen|skip to content|"
    r"menü schließen|newsletter abonnieren)", re.I)
ERROR_PREFIX = "__ERROR__"


# ----------------------------------------------------------------------------
# Helpers
# ----------------------------------------------------------------------------
def load_config(path: Path) -> dict:
    with open(path, encoding="utf-8") as f:
        return yaml.safe_load(f)


def log(*args):
    print(*args, file=sys.stderr, flush=True)


def load_recipients(path: Path, cfg: dict) -> list[dict]:
    """Read CSV (comma/semicolon/tab) or XLSX and map columns."""
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        ws = load_workbook(path, read_only=True, data_only=True).active
        rows = list(ws.iter_rows(values_only=True))
        header = [str(h or "").strip() for h in rows[0]]
        records = [{header[i]: ("" if v is None else str(v).strip()) for i, v in enumerate(row)}
                   for row in rows[1:] if any(row)]
    else:
        raw = path.read_text(encoding="utf-8-sig")
        dialect = csv.Sniffer().sniff(raw[:4096], delimiters=",;\t")
        records = [{k.strip(): (v or "").strip() for k, v in r.items() if k}
                   for r in csv.DictReader(raw.splitlines(), dialect=dialect)]

    if not records:
        sys.exit("Recipient list is empty.")
    mapping = {}
    header_lower = {k.lower(): k for k in records[0]}
    for field, aliases in cfg["columns"].items():
        for alias in aliases:
            if alias.lower() in header_lower:
                mapping[field] = header_lower[alias.lower()]
                break
    if "url" not in mapping:
        sys.exit(f"No URL column found. Columns: {list(records[0])}")

    # Internal fields are prefixed with "_" so they never collide with input columns.
    for i, r in enumerate(records, 1):
        r["_row"] = i
        r["_url"] = r.get(mapping["url"], "")
        r["_agency"] = r.get(mapping.get("agency", ""), "")
        r["_first_name"] = r.get(mapping.get("first_name", ""), "")
    return [r for r in records if r["_url"]]


def normalize_url(url: str) -> str:
    url = url.strip()
    if not re.match(r"^https?://", url, re.I):
        url = "https://" + url
    return url.rstrip("/")


def domain_of(url: str) -> str:
    return urlparse(url).netloc.lower().removeprefix("www.")


def cache_path(url: str) -> Path:
    d = domain_of(normalize_url(url)) or hashlib.md5(url.encode()).hexdigest()
    return CACHE_DIR / (re.sub(r"[^a-z0-9.-]", "_", d) + ".json")


def estimate_tokens(text: str) -> int:
    return int(len(text) / 3.2) + 1  # German text, newer tokenizer: on the safe side


# ----------------------------------------------------------------------------
# 1) Crawling
# ----------------------------------------------------------------------------
def fetch(session: requests.Session, url: str, timeout: int) -> tuple[str, str] | None:
    """Return (final_url, html) or None on any failure."""
    try:
        r = session.get(url, timeout=timeout, allow_redirects=True)
        if r.status_code >= 400 or "html" not in r.headers.get("content-type", "html"):
            return None
        if not r.encoding or r.encoding.lower() == "iso-8859-1":
            r.encoding = r.apparent_encoding
        return r.url, r.text
    except requests.RequestException:
        return None


def find_subpages(html: str, base_url: str, cfg: dict) -> list[str]:
    """Score internal links by keywords and return the best ones."""
    ccfg = cfg["crawl"]
    soup = BeautifulSoup(html, "html.parser")
    own_domain = domain_of(base_url)
    candidates: dict[str, int] = {}

    for a in soup.find_all("a", href=True):
        href = urldefrag(urljoin(base_url, a["href"]))[0].rstrip("/")
        parsed = urlparse(href)
        if parsed.scheme not in ("http", "https") or domain_of(href) != own_domain:
            continue
        path = parsed.path.lower()
        if not path or path == "/" or path.endswith(FILE_EXTENSIONS):
            continue
        haystack = (a.get_text(" ", strip=True) + " " + path).lower()
        if any(x in haystack for x in ccfg["exclude"]):
            continue
        score = 0
        for weight, keywords in ccfg["keywords"].items():
            if any(k in haystack for k in keywords):
                score = max(score, int(weight))
        if score == 0:
            continue
        depth = path.strip("/").count("/")
        score -= depth * 3              # deep paths are usually blog posts
        if a.find_parent(["nav", "header"]):
            score += 2                  # main navigation = important pages
        candidates[href] = max(candidates.get(href, -99), score)

    ranked = sorted(candidates.items(), key=lambda kv: -kv[1])
    return [u for u, s in ranked if s > 0][: ccfg["max_subpages"]]


def clean_html(html: str) -> str:
    """HTML -> compact plain text, with title and meta description prepended."""
    soup = BeautifulSoup(html, "html.parser")
    head = []
    if soup.title and soup.title.string:
        head.append("Titel: " + soup.title.string.strip())
    meta = soup.find("meta", attrs={"name": "description"})
    if meta and meta.get("content"):
        head.append("Beschreibung: " + meta["content"].strip())

    text = ""
    if trafilatura:
        text = trafilatura.extract(html, include_comments=False, include_tables=True,
                                   include_links=False, favor_recall=True, deduplicate=True) or ""
    if len(text) < 200:  # fallback, e.g. pages built from many tiny blocks
        for tag in soup(["script", "style", "noscript", "svg", "iframe", "form",
                         "nav", "footer", "header", "aside", "button"]):
            tag.decompose()
        for tag in soup.select('[class*="cookie"], [id*="cookie"], [class*="consent"], [id*="consent"]'):
            tag.decompose()
        text = soup.get_text("\n", strip=True)

    return "\n".join(head + [text])


def compact(pages: list[dict], cfg: dict) -> str:
    """Normalize lines, drop cross-page duplicates and noise, enforce length limits."""
    ccfg = cfg["crawl"]
    seen: set[str] = set()
    blocks, total = [], 0
    for page in pages:
        lines = []
        for line in page["text"].splitlines():
            line = re.sub(r"\s+", " ", line).strip()
            key = line.lower()
            if len(line) < 3 or key in seen or (len(line) < 120 and NOISE.search(line)):
                continue
            seen.add(key)
            lines.append(line)
        text = "\n".join(lines)[: ccfg["max_chars_per_page"]]
        remaining = ccfg["max_chars_total"] - total
        if remaining <= 200:
            break
        text = text[:remaining]
        total += len(text)
        blocks.append(f"### Seite: {urlparse(page['url']).path or '/'}\n{text}")
    return "\n\n".join(blocks)


def crawl_site(url: str, cfg: dict, refresh: bool = False) -> dict:
    """Crawl one site (or load from cache). Failures are cached too -> use --refresh."""
    cp = cache_path(url)
    if cp.exists() and not refresh:
        return json.loads(cp.read_text(encoding="utf-8"))

    ccfg = cfg["crawl"]
    session = requests.Session()
    session.headers.update({"User-Agent": USER_AGENT, "Accept-Language": "de-DE,de;q=0.9"})
    start = normalize_url(url)
    attempts = [start]
    if "://www." not in start:
        attempts.append(start.replace("://", "://www.", 1))
    attempts.append(start.replace("https://", "http://", 1))

    result = {"url": url, "status": "error", "pages": [], "text": "",
              "fetched_at": datetime.now().isoformat(timespec="seconds")}
    hit = None
    for attempt in attempts:
        hit = fetch(session, attempt, ccfg["timeout"])
        if hit:
            break
    if not hit:
        result["note"] = "Homepage not reachable"
    else:
        final_url, html = hit
        pages = [{"url": final_url, "text": clean_html(html)}]
        for sub in find_subpages(html, final_url, cfg):
            fetched = fetch(session, sub, ccfg["timeout"])
            if fetched:
                pages.append({"url": fetched[0], "text": clean_html(fetched[1])})
        result["pages"] = [p["url"] for p in pages]
        result["text"] = compact(pages, cfg)
        result["status"] = "ok" if len(result["text"]) > 300 else "thin"
        if result["status"] == "thin":
            result["note"] = "Very little text (JS-rendered site or under construction?)"

    CACHE_DIR.mkdir(parents=True, exist_ok=True)
    cp.write_text(json.dumps(result, ensure_ascii=False, indent=1), encoding="utf-8")
    return result


def crawl_all(recipients: list[dict], cfg: dict, refresh: bool) -> dict[int, dict]:
    sites = {}
    with ThreadPoolExecutor(cfg["crawl"]["workers"]) as ex:
        futures = {ex.submit(crawl_site, r["_url"], cfg, refresh): r for r in recipients}
        for i, fut in enumerate(as_completed(futures), 1):
            r = futures[fut]
            sites[r["_row"]] = site = fut.result()
            log(f"  [{i}/{len(recipients)}] {site['status']:5} {len(site['pages'])} pages, "
                f"~{estimate_tokens(site['text'])} tokens  {r['_url']}")
    return sites


# ----------------------------------------------------------------------------
# 2) Prompt + model backends
# ----------------------------------------------------------------------------
def build_user_message(r: dict, site: dict) -> str:
    # Message content stays German: it is read by the model, next to a German prompt.
    parts = []
    if r["_agency"]:
        parts.append(f"Agentur laut unserer Liste: {r['_agency']}")
    parts.append(f"URL: {r['_url']}")
    parts.append(f"<website>\n{site['text'] or '(kein Text gefunden)'}\n</website>")
    return "\n".join(parts)


class Answer(NamedTuple):
    text: str
    tokens_in: int = 0
    tokens_out: int = 0
    seconds: float = 0.0
    truncated: bool = False   # model hit max_tokens (thinking counts toward it!)


def error_answer(msg: str, seconds: float = 0.0) -> Answer:
    return Answer(f"{ERROR_PREFIX} {msg}", seconds=seconds)


class AnthropicBackend:
    def __init__(self, profile: dict):
        import anthropic
        if not os.environ.get("ANTHROPIC_API_KEY"):
            sys.exit("ANTHROPIC_API_KEY is not set (export ANTHROPIC_API_KEY=...).")
        # Keys that are not scoped to a workspace need the workspace ID as a header.
        workspace = profile.get("workspace_id") or os.environ.get("ANTHROPIC_WORKSPACE_ID")
        headers = {"anthropic-workspace-id": workspace} if workspace else None
        self.client = anthropic.Anthropic(max_retries=6, default_headers=headers)
        self.profile = profile

    def _params(self, system: str, user: str) -> dict:
        p = self.profile
        params = dict(model=p["model"], max_tokens=p.get("max_tokens", 8000),
                      system=[{"type": "text", "text": system, "cache_control": {"type": "ephemeral"}}],
                      messages=[{"role": "user", "content": user}])
        if "effort" in p:
            params["output_config"] = {"effort": p["effort"]}
        if "temperature" in p:
            params["temperature"] = p["temperature"]
        params.update(p.get("extra", {}))
        return params

    @staticmethod
    def _to_answer(msg) -> Answer:
        text = "".join(b.text for b in msg.content if b.type == "text")
        u = msg.usage
        tokens_in = u.input_tokens + (u.cache_read_input_tokens or 0) + (u.cache_creation_input_tokens or 0)
        return Answer(text, tokens_in, u.output_tokens, truncated=msg.stop_reason == "max_tokens")

    def ask(self, system: str, user: str) -> Answer:
        return self._to_answer(self.client.messages.create(**self._params(system, user)))

    def ask_batch(self, jobs: list[tuple[str, str, str]]) -> dict[str, Answer]:
        """jobs: (custom_id, system, user). Blocks until the batch has finished."""
        batch = self.client.messages.batches.create(requests=[
            {"custom_id": cid, "params": self._params(system, user)} for cid, system, user in jobs])
        log(f"  Batch {batch.id} submitted ({len(jobs)} requests). Waiting ...")
        while True:
            batch = self.client.messages.batches.retrieve(batch.id)
            if batch.processing_status == "ended":
                break
            c = batch.request_counts
            log(f"    running: {c.processing} pending, {c.succeeded} done, {c.errored} errors")
            time.sleep(30)
        out = {}
        for entry in self.client.messages.batches.results(batch.id):
            if entry.result.type == "succeeded":
                out[entry.custom_id] = self._to_answer(entry.result.message)
            else:
                out[entry.custom_id] = error_answer(entry.result.type)
        return out


class OpenAICompatibleBackend:
    """LM Studio, Ollama, llama.cpp server, vLLM ... anything with /v1/chat/completions."""
    def __init__(self, profile: dict):
        self.profile = profile
        self.url = profile["base_url"].rstrip("/") + "/chat/completions"

    def ask(self, system: str, user: str) -> Answer:
        p = self.profile
        body = {"model": p["model"], "max_tokens": p.get("max_tokens", 1500),
                "messages": [{"role": "system", "content": system},
                             {"role": "user", "content": user}]}
        if "temperature" in p:
            body["temperature"] = p["temperature"]
        body.update(p.get("extra", {}))
        r = requests.post(self.url, json=body, timeout=p.get("timeout", 300),
                          headers={"Authorization": "Bearer local"})
        r.raise_for_status()
        data = r.json()
        usage = data.get("usage") or {}
        choice = data["choices"][0]
        return Answer(choice["message"].get("content") or "",
                      usage.get("prompt_tokens", 0), usage.get("completion_tokens", 0),
                      truncated=choice.get("finish_reason") == "length")


BACKENDS = {"anthropic": AnthropicBackend, "openai": OpenAICompatibleBackend}


# ----------------------------------------------------------------------------
# 3) Response parsing + validation
# ----------------------------------------------------------------------------
def parse_json(text: str) -> dict | None:
    text = re.sub(r"<think>.*?</think>", "", text, flags=re.S)  # reasoning of local models
    text = re.sub(r"^```(?:json)?|```$", "", text.strip(), flags=re.M).strip()
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", text, re.S)
        if match:
            try:
                return json.loads(match.group(0))
            except json.JSONDecodeError:
                return None
    return None


def count_words(text: str) -> int:
    """Count tokens containing a letter or digit (dashes are not words)."""
    return sum(1 for tok in text.split() if re.search(r"\w", tok))


def validate(data: dict | None, site: dict, cfg: dict) -> list[str]:
    """Return a list of flags. Never modifies the text."""
    rules = cfg["validation"]
    flags = []
    if site["status"] != "ok":
        flags.append(f"crawl_{site['status']}")
    if data is None:
        return flags + ["no_json"]
    paragraph = data.get("absatz") or ""
    if not paragraph:
        return flags + ["no_paragraph"]
    words = count_words(paragraph)
    if words > rules["max_words"]:
        flags.append(f"too_long_{words}_words")
    sentences = len(re.findall(r"[.?]+(\s|$)", paragraph))  # rough: "z. B." counts as a sentence end
    if sentences > rules["max_sentences"]:
        flags.append(f"{sentences}_sentences")
    for ch in rules["forbidden_chars"]:
        if ch in paragraph:
            flags.append(f"char_{ch}")
    for phrase in rules["forbidden_phrases"]:
        if phrase.lower() in paragraph.lower():
            flags.append(f"phrase_{phrase.replace(' ', '_')}")
    name = (data.get("agenturname") or "").strip()
    if name and name.lower() not in paragraph.lower():
        flags.append("name_missing")
    if (data.get("sicherheit") or "").lower().startswith("niedrig"):
        flags.append("low_confidence")
    return flags


# ----------------------------------------------------------------------------
# 3b) Assembly mode (config has "assembly"): model returns evidence, we build the paragraph
# ----------------------------------------------------------------------------
QUOTE_CHARS = str.maketrans({c: "" for c in "„“”\"'‚‘’«»"})


def normalize_text(text: str) -> str:
    """Lowercase, drop quote marks, unify dashes/spaces: for 'does this quote appear on the site'."""
    text = (text or "").translate(QUOTE_CHARS).replace("\u00a0", " ").lower()
    text = re.sub(r"[–—]", "-", text)
    return re.sub(r"\s+", " ", text).strip()


def numbers_in(text: str) -> set[str]:
    """Numbers as digit strings: '3.000' -> '3000', '1,5' -> '15', '3 000' -> '3000'."""
    found = re.findall(r"\d+(?:[.,]\d+)*(?:[ \u00a0]\d{3}\b)*", text or "")
    return {re.sub(r"\D", "", n) for n in found}


def assemble(d: dict, r: dict, site: dict, cfg: dict) -> tuple[dict, list[str]]:
    """Check the model's evidence against the site text and build the paragraph.
    Anything that cannot be verified falls back to the safe variant and gets a flag."""
    acfg = cfg["assembly"]
    flags: list[str] = []
    site_text = site.get("text", "")
    site_norm = normalize_text(site_text)
    site_numbers = numbers_in(site_text)

    kind = (d.get("typ") or "keiner").strip().lower()
    evidence = re.sub(r"\s+", " ", (d.get("beleg") or "")).strip().rstrip(".")
    quote = (d.get("zitat") or "").strip()

    if kind in ("ergebnis", "system"):
        starts = acfg["evidence_start"]["result" if kind == "ergebnis" else "system"]
        problem = None
        if not evidence:
            problem = "evidence_missing"
        elif evidence.split()[0].lower() not in starts:
            problem = "evidence_grammar"
        elif count_words(evidence) > acfg.get("max_evidence_words", 14):
            problem = "evidence_too_long"
        elif any(ch in evidence for ch in cfg["validation"]["forbidden_chars"] + ['"', "„", "“"]):
            problem = "evidence_chars"
        else:
            missing = numbers_in(evidence) - site_numbers
            if missing:
                problem = "number_not_on_site_" + "_".join(sorted(missing))
        if problem:
            flags.append(f"{problem}->fallback")
            kind = "keiner"
        elif quote and normalize_text(quote) not in site_norm:
            flags.append("quote_not_found")
        elif not quote:
            flags.append("quote_missing")
    if kind not in ("ergebnis", "system"):
        kind, evidence = "keiner", ""

    # Decision maker: industry term only when certain and backed by a quote from the site
    makers = acfg["decision_makers"]
    industry = (d.get("branche") or "").strip().lower()
    industry_quote = (d.get("branche_zitat") or "").strip()
    used_industry = "default"
    if industry in makers and industry not in ("default", "geschaeftsfuehrer"):
        sure = (d.get("branche_sicherheit") or "").lower().startswith("hoch")
        if sure and industry_quote and normalize_text(industry_quote) in site_norm:
            used_industry = industry
        elif sure:
            flags.append(f"industry_{industry}_quote_not_found->default")

    if kind == "ergebnis":
        template = acfg["templates"]["result_numbers" if re.search(r"\d", evidence) else "result_plain"]
    elif kind == "system":
        template = acfg["templates"]["system"]
    else:
        template = acfg["templates"]["none"]
    evidence = evidence[:1].lower() + evidence[1:]
    paragraph = template.format(name=r["_agency"], who=makers[used_industry], evidence=evidence)
    info = {"paragraph": paragraph, "type": kind, "evidence": evidence, "quote": quote,
            "source": d.get("quelle") or "", "industry": used_industry,
            "industry_model": industry, "industry_quote": industry_quote}
    return info, flags


# ----------------------------------------------------------------------------
# 4) Commands
# ----------------------------------------------------------------------------
def select_recipients(recipients, args):
    if args.rows:
        wanted = {int(x) for x in args.rows.split(",")}
        recipients = [r for r in recipients if r["_row"] in wanted]
    if args.sample:
        random.seed(args.seed)
        recipients = random.sample(recipients, min(args.sample, len(recipients)))
    if args.limit:
        recipients = recipients[: args.limit]
    return recipients


def cmd_crawl(args, cfg):
    recipients = select_recipients(load_recipients(Path(args.list), cfg), args)
    log(f"Crawling {len(recipients)} websites ...")
    sites = crawl_all(recipients, cfg, args.refresh)
    tokens = [estimate_tokens(s["text"]) for s in sites.values()]
    ok = sum(s["status"] == "ok" for s in sites.values())
    log(f"\nDone: {ok}/{len(sites)} ok, avg ~{sum(tokens) // max(len(tokens), 1)} tokens of website text.")
    log(f"Cache: {CACHE_DIR}")


def cmd_prompt(args, cfg):
    recipients = load_recipients(Path(args.list), cfg)
    r = next((x for x in recipients if x["_row"] == args.row), None) or sys.exit("Row not found.")
    site = crawl_site(r["_url"], cfg, args.refresh)
    system = (BASE_DIR / args.prompt).read_text(encoding="utf-8")
    user = build_user_message(r, site)
    print("=== SYSTEM ===\n" + system + "\n=== USER ===\n" + user)
    log(f"\n~{estimate_tokens(system)} + ~{estimate_tokens(user)} input tokens")


def query_profile(name: str, profile: dict, system: str, jobs: list, use_batch: bool) -> dict[int, Answer]:
    """Run all jobs (recipient, site, user_msg) against one profile. Returns {row: Answer}."""
    backend = BACKENDS[profile["backend"]](profile)
    answers: dict[int, Answer] = {}

    if use_batch and profile["backend"] == "anthropic" and jobs:
        t0 = time.time()
        res = backend.ask_batch([(f"r{r['_row']}", system, user) for r, _, user in jobs])
        secs = (time.time() - t0) / len(jobs)
        for r, _, _ in jobs:
            answers[r["_row"]] = res.get(f"r{r['_row']}", error_answer("missing"))._replace(seconds=secs)
        return answers

    def one(r, user) -> tuple[int, Answer]:
        t0 = time.time()
        try:
            return r["_row"], backend.ask(system, user)._replace(seconds=time.time() - t0)
        except Exception as e:  # noqa: BLE001 - one failure must not kill the run
            return r["_row"], error_answer(str(e), time.time() - t0)

    first_errors: list[str] = []
    aborted = None
    with ThreadPoolExecutor(max(profile.get("workers", 1), 1)) as ex:
        futures = [ex.submit(one, r, user) for r, _, user in jobs]
        for i, fut in enumerate(as_completed(futures), 1):
            row, ans = fut.result()
            answers[row] = ans
            err = ans.text.startswith(ERROR_PREFIX)
            status = (f"  ERROR: {ans.text[len(ERROR_PREFIX):].strip()[:200]}" if err
                      else f"  TRUNCATED at {ans.tokens_out} tokens (raise max_tokens)" if ans.truncated
                      else f"  {ans.tokens_out} tokens out")
            log(f"  [{i}/{len(futures)}] row {row}: {ans.seconds:.1f}s" + status)
            # Fail fast: the first 3 results are identical errors -> config/auth problem, stop.
            if i <= 3:
                first_errors.append(ans.text if err else "")
                if i == 3 and all(first_errors) and len(set(first_errors)) == 1:
                    aborted = first_errors[0]
                    log(f"  Aborting profile '{name}': first 3 requests failed with the same error.")
                    ex.shutdown(wait=True, cancel_futures=True)
                    break
    if aborted:
        for r, _, _ in jobs:
            answers.setdefault(r["_row"], Answer(aborted))
    return answers


def cmd_run(args, cfg):
    profiles = list(dict.fromkeys(args.profile or cfg.get("default_profiles", ["sonnet"])))
    for p in profiles:
        if p not in cfg["profiles"]:
            sys.exit(f"Unknown profile '{p}'. Available: {', '.join(cfg['profiles'])}")

    recipients = select_recipients(load_recipients(Path(args.list), cfg), args)
    system = (BASE_DIR / args.prompt).read_text(encoding="utf-8")
    run_name = args.name or "-".join(profiles)
    out_dir = RUNS_DIR / f"{datetime.now():%Y-%m-%d_%H%M}_{re.sub(r'[^A-Za-z0-9_-]', '_', run_name)}"
    out_dir.mkdir(parents=True, exist_ok=True)

    log(f"Step 1/2: loading {len(recipients)} websites (cache is used) ...")
    sites = crawl_all(recipients, cfg, args.refresh)

    all_jobs = [(r, sites[r["_row"]], build_user_message(r, sites[r["_row"]])) for r in recipients]
    jobs = [j for j in all_jobs if j[1]["status"] != "error"]  # unreachable sites never hit a model
    if len(jobs) < len(all_jobs):
        log(f"  {len(all_jobs) - len(jobs)} website(s) unreachable, skipped")

    results = []
    for pname in profiles:
        profile = cfg["profiles"][pname]
        log(f"\nStep 2/2: profile '{pname}' ({profile['model']}) ...")
        answers = query_profile(pname, profile, system, jobs, args.batch)
        discount = 0.5 if (args.batch and profile["backend"] == "anthropic") else 1.0

        for r, site, _ in all_jobs:
            ans = answers.get(r["_row"], error_answer("website unreachable, not sent to model"))
            raw, tok_in, tok_out, secs = ans.text, ans.tokens_in, ans.tokens_out, ans.seconds
            data = None if raw.startswith(ERROR_PREFIX) else parse_json(raw)
            d = data or {}
            info, extra_flags = {}, []
            if cfg.get("assembly") and data is not None:
                info, extra_flags = assemble(d, r, site, cfg)
                d["absatz"], d["agenturname"] = info["paragraph"], r["_agency"]
            flags = validate(data, site, cfg) + extra_flags
            if ans.truncated:
                flags.insert(0, "truncated")
            cost = (tok_in * profile.get("price_input", 0) + tok_out * profile.get("price_output", 0)) / 1e6
            results.append({
                "row": r["_row"], "profile": pname, "model": profile["model"],
                "agency_list": r["_agency"], "first_name": r["_first_name"], "url": r["_url"],
                "agency_name": d.get("agenturname", ""),
                "angle": d.get("winkel", ""),
                "industries": ", ".join(d.get("zielbranchen") or []),
                "roles": ", ".join(d.get("engpassberufe") or []),
                "paragraph": d.get("absatz") or "",
                "alternative": d.get("alternative") or "",
                "words": count_words(d.get("absatz") or ""),
                "confidence": d.get("sicherheit", ""),
                "flags": " ".join(flags),
                "evidence": " | ".join(d.get("belege") or []),
                "note": d.get("hinweis") or site.get("note", "") or (raw[:300] if data is None else ""),
                "tokens_in": tok_in, "tokens_out": tok_out,
                "cost_usd": round(cost * discount, 5), "seconds": round(secs, 1),
                "pages": " ".join(site["pages"]),
                # assembly mode only (empty otherwise)
                "c_type": info.get("type", ""), "c_evidence": info.get("evidence", ""),
                "c_quote": info.get("quote", ""), "c_source": info.get("source", ""),
                "c_industry": info.get("industry", ""), "c_industry_model": info.get("industry_model", ""),
                "c_industry_quote": info.get("industry_quote", ""),
                "_raw": raw, "_record": r,
            })

    write_outputs(out_dir, results, profiles, recipients, args)


def write_outputs(out_dir: Path, results: list[dict], profiles: list[str], recipients, args):
    # Raw answers, for debugging
    with open(out_dir / "raw_responses.jsonl", "w", encoding="utf-8") as f:
        for e in results:
            f.write(json.dumps({k: v for k, v in e.items() if k != "_record"}, ensure_ascii=False) + "\n")

    # Long table: one line per recipient x profile (semicolon -> opens cleanly in German Excel)
    fields = [k for k in results[0] if not k.startswith("_")]
    with open(out_dir / "results.csv", "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields, delimiter=";", extrasaction="ignore")
        w.writeheader()
        w.writerows(sorted(results, key=lambda e: (e["row"], e["profile"])))

    # Wide table: models side by side (only when comparing)
    if len(profiles) > 1:
        idx = {(e["row"], e["profile"]): e for e in results}
        with open(out_dir / "comparison.csv", "w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f, delimiter=";")
            w.writerow(["row", "agency", "url"] + [c for p in profiles for c in (f"paragraph_{p}", f"flags_{p}")])
            for r in recipients:
                w.writerow([r["_row"], r["_agency"], r["_url"]] +
                           [v for p in profiles for v in (idx[(r["_row"], p)]["paragraph"],
                                                           idx[(r["_row"], p)]["flags"])])

    # InDesign data merge: original columns + paragraph (UTF-8, comma)
    for p in profiles:
        entries = sorted((e for e in results if e["profile"] == p), key=lambda e: e["row"])
        orig_cols = [k for k in entries[0]["_record"]
                     if not k.startswith("_") and k not in ("AgencyName", "Paragraph", "Review")]
        with open(out_dir / f"datamerge_{p}.csv", "w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f)
            w.writerow(orig_cols + ["AgencyName", "Paragraph", "Review"])
            for e in entries:
                w.writerow([e["_record"].get(k, "") for k in orig_cols] +
                           [e["agency_name"] or e["agency_list"], e["paragraph"], "yes" if e["flags"] else ""])

    # Summary
    summary = {"created_at": datetime.now().isoformat(timespec="seconds"), "args": vars(args), "profiles": {}}
    log("\n" + "=" * 70)
    for p in profiles:
        es = [e for e in results if e["profile"] == p]
        n = max(len(es), 1)
        clean = sum(1 for e in es if e["paragraph"] and not e["flags"])
        cost = sum(e["cost_usd"] for e in es)
        s = summary["profiles"][p] = {
            "recipients": len(es), "clean": clean, "cost_usd": round(cost, 4),
            "cost_per_recipient": round(cost / n, 5),
            "seconds_per_recipient": round(sum(e["seconds"] for e in es) / n, 1),
            "avg_tokens_in": sum(e["tokens_in"] for e in es) // n,
            "avg_tokens_out": sum(e["tokens_out"] for e in es) // n}
        log(f"{p:16} {clean:>3}/{len(es)} clean | ${s['cost_usd']:.4f} total, "
            f"${s['cost_per_recipient']:.4f}/recipient | {s['seconds_per_recipient']}s/recipient | "
            f"avg {s['avg_tokens_in']} in / {s['avg_tokens_out']} out")
    (out_dir / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    log("=" * 70 + f"\nOutput: {out_dir}")


# ----------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--config", default="config.yaml")
    sub = ap.add_subparsers(dest="cmd", required=True)

    def add_common(sp):
        sp.add_argument("list", help="CSV or XLSX with at least a URL column")
        sp.add_argument("--refresh", action="store_true", help="re-crawl instead of using the cache")

    def add_filters(sp):
        sp.add_argument("--limit", type=int, help="only the first N recipients")
        sp.add_argument("--sample", type=int, help="N random recipients")
        sp.add_argument("--seed", type=int, default=42, help="random seed for --sample")
        sp.add_argument("--rows", help="specific row numbers, e.g. 1,4,7")

    c = sub.add_parser("crawl", help="only crawl and cache websites")
    add_common(c); add_filters(c)

    r = sub.add_parser("run", help="generate paragraphs")
    add_common(r); add_filters(r)
    r.add_argument("--profile", action="append", help="profile from config.yaml, repeatable")
    r.add_argument("--batch", action="store_true", help="Claude Batch API (50%% cheaper, async)")
    r.add_argument("--prompt", default="prompt.md", help="prompt file (for prompt variants)")
    r.add_argument("--name", help="run name (used for the output folder)")

    p = sub.add_parser("prompt", help="print the full prompt for one recipient")
    add_common(p)
    p.add_argument("--row", type=int, default=1)
    p.add_argument("--prompt", default="prompt.md")

    args = ap.parse_args()
    cfg = load_config(BASE_DIR / args.config)
    {"crawl": cmd_crawl, "run": cmd_run, "prompt": cmd_prompt}[args.cmd](args, cfg)


if __name__ == "__main__":
    main()
